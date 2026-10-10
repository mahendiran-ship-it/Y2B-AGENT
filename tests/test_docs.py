"""Offline document generation tests.   Run with:  python -m unittest discover -s tests -v

Every test writes into temporary folders. Android storage is simulated by pointing y2b_docs.SHARED_ROOTS at a
temp dir, so nothing here ever touches a real /storage/emulated/0.
"""
import builtins
import importlib.util
import io
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, ".."))
sys.path.insert(0, HERE)
os.environ.setdefault("Y2B_HOME", tempfile.mkdtemp())

import y2b_agent as y2b  # noqa: E402
import y2b_docs as D  # noqa: E402
import fake_server  # noqa: E402

HAVE = {f: not D.missing_libraries([f]) for f in D.FORMATS}
ALL_LIBS = all(HAVE.values())


def pdf_text(path):
    """Extract text from a PDF with whatever parser is installed (None if there is none)."""
    if importlib.util.find_spec("pypdf"):
        import pypdf
        return "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(str(path)).pages)
    if shutil.which("pdftotext"):
        return subprocess.run(["pdftotext", str(path), "-"], capture_output=True, text=True).stdout
    return None


REPORT = """# Network Report
An intro with **bold**, *italic*, café and ₹100.

## Findings
- First point
  - Nested point
- Second point

1. Step one
2. Step two

## Data
| Name | Score |
|---|---|
| Asha | 90 |
| Ravi | 78 |

[pagebreak]
## Conclusion
All done.
"""

DECK = """# Marine Pollution
Protecting our oceans

## Causes
- Plastic waste
- Oil spills

## Steps
1. Reduce
2. Recycle

## Data
| Year | Tonnes |
|---|---|
| 2020 | 8 |
| 2021 | 9 |
"""

SHEET = """# Marks
## Students
| Name | Maths | Science | Total |
|---|---|---|---|
| Asha | 90 | 85 | =SUM(B{row}:C{row}) |
| Ravi | 78 | 88 | =SUM(B{row}:C{row}) |
## Notes
| Key | Value |
|---|---|
| Pass mark | 40 |
"""


class TmpCase(unittest.TestCase):
    """Gives each test a workspace, a fake 'Android storage' root and restores SHARED_ROOTS afterwards."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp()).resolve()
        self.addCleanup(shutil.rmtree, str(self.tmp), True)
        self.work = self.tmp / "work"
        self.work.mkdir()
        self.shared = self.tmp / "storage" / "emulated" / "0"
        self.shared.mkdir(parents=True)
        old = list(D.SHARED_ROOTS)
        D.SHARED_ROOTS[:] = [str(self.shared)]
        self.addCleanup(lambda: D.SHARED_ROOTS.__setitem__(slice(None), old))
        self.paths = D.DocPaths(self.work)

    def save(self, fmt, md, name="doc", **kw):
        return D.save_document(fmt, md, name, self.work, self.paths, **kw)


# ═════════════════════════ markdown parsing ═════════════════════════
class MarkdownTests(unittest.TestCase):
    def test_block_types(self):
        b = D.parse_markdown(REPORT)
        self.assertEqual([x["type"] for x in b],
                         ["title", "paragraph", "heading", "bullets", "numbered", "heading", "table",
                          "pagebreak", "heading", "paragraph"])
        self.assertEqual(b[0]["text"], "Network Report")
        self.assertEqual(b[3]["items"], [("First point", 0), ("Nested point", 1), ("Second point", 0)])
        self.assertEqual(b[6]["header"], ["Name", "Score"])
        self.assertEqual(b[6]["rows"], [["Asha", "90"], ["Ravi", "78"]])

    def test_inline_formatting(self):
        self.assertEqual(D.parse_inline("a **b** *c* `d`"),
                         [("a ", False, False), ("b", True, False), (" ", False, False), ("c", False, True),
                          (" ", False, False), ("d", False, False)])
        self.assertEqual(D.plain("**x** and *y*"), "x and y")

    def test_fences_and_control_characters_are_removed(self):
        b = D.parse_markdown("```markdown\n# T\nhello\x00 world\n```")
        self.assertEqual(b[0], {"type": "title", "text": "T"})
        self.assertEqual(b[1]["text"], "hello world")

    def test_images_are_dropped_with_a_warning(self):
        w = []
        b = D.parse_markdown("# T\n![x](http://e.com/a.png)\ntext", w)
        self.assertEqual(D.all_text(b), "T\ntext")
        self.assertTrue(any("image" in x for x in w))

    def test_csv_fallback_for_spreadsheets(self):
        b = D.csv_to_blocks("Name,Marks\nAsha,90\nRavi,78\n")
        self.assertEqual(b[0]["header"], ["Name", "Marks"])
        self.assertEqual(len(b[0]["rows"]), 2)

    def test_empty_content_is_rejected(self):
        for fmt in D.FORMATS:
            with self.assertRaises(D.DocError):
                D.prepare_blocks(fmt, "   \n\n  ")


# ═════════════════════════ the four formats ═════════════════════════
@unittest.skipUnless(HAVE["pdf"], "reportlab not installed")
class PdfTests(TmpCase):
    def test_pdf_is_valid_and_has_the_text(self):
        r = self.save("pdf", REPORT, "report")
        self.assertEqual(r.path.name, "report.pdf")
        self.assertGreater(r.size, 0)
        self.assertEqual(r.details["pages"], 2)  # the [pagebreak] forces a second page
        self.assertTrue(r.path.read_bytes().startswith(b"%PDF-"))
        text = pdf_text(r.path)
        if text is not None:
            for needle in ("Network Report", "Findings", "Nested point", "Asha", "Conclusion"):
                self.assertIn(needle, text)

    def test_long_text_wraps_and_paginates(self):
        md = "# Long\n" + "\n\n".join("Paragraph %d " % i + "lorem ipsum dolor sit amet " * 25 for i in range(60))
        md += "\n\n" + "x" * 400  # one unbreakable word must not crash or overflow
        r = self.save("pdf", md, "long")
        self.assertGreaterEqual(r.details["pages"], 4)
        text = pdf_text(r.path)
        if text is not None:
            self.assertIn("Paragraph 59", text)

    def test_unicode_is_drawn_or_reported(self):
        r = self.save("pdf", "# T\ncafé ₹100 “quoted” 你好", "uni")
        text = pdf_text(r.path)
        fam, has = D.pdf_font()
        if text is not None:
            self.assertIn("caf", text)
        if has is None:  # no Unicode font on this machine: the unsupported characters are reported, not silently lost
            self.assertTrue(any("Unicode font" in w for w in r.warnings))

    def test_without_a_unicode_font_unsupported_characters_become_question_marks_with_a_warning(self):
        # the usual situation on a fresh Termux install: only the built-in Helvetica is available
        with mock.patch.dict(D._FONT_CACHE, {"font": ("Helvetica", None)}):
            r = self.save("pdf", "# T\ncaf\u00e9 \u20ac5 \u201cq\u201d \u20b9100 \u4f60\u597d", "nofont")
        self.assertTrue(r.path.is_file())
        self.assertEqual(len([w for w in r.warnings if "Unicode font" in w]), 1)
        self.assertIn("3 character", r.warnings[0])  # the rupee sign and 2 Chinese characters; the euro sign and quotes are fine
        text = pdf_text(r.path)
        if text is not None:
            self.assertIn("caf\u00e9", text)
            self.assertIn("5", text)
            self.assertNotIn("\u20b9", text)

    def test_markup_characters_in_text_do_not_break_the_pdf(self):
        r = self.save("pdf", "# A & B <tag>\nUse <b>this</b> & that </para> 5 < 6", "esc")
        text = pdf_text(r.path)
        if text is not None:
            self.assertIn("A & B <tag>", text)


@unittest.skipUnless(HAVE["docx"], "python-docx not installed")
class DocxTests(TmpCase):
    def test_docx_reopens_with_expected_content(self):
        import docx
        r = self.save("docx", REPORT, "report")
        d = docx.Document(str(r.path))
        by_style = {}
        for p in d.paragraphs:
            by_style.setdefault(p.style.name, []).append(p.text)
        self.assertEqual(by_style["Title"], ["Network Report"])
        self.assertEqual(by_style["Heading 1"], ["Findings", "Data", "Conclusion"])
        self.assertEqual(by_style["List Bullet"], ["First point", "Second point"])
        self.assertEqual(by_style["List Bullet 2"], ["Nested point"])
        self.assertEqual(by_style["List Number"], ["Step one", "Step two"])
        self.assertEqual([[c.text for c in row.cells] for row in d.tables[0].rows],
                         [["Name", "Score"], ["Asha", "90"], ["Ravi", "78"]])
        intro = [p for p in d.paragraphs if p.text.startswith("An intro")][0]
        self.assertIn("café", intro.text)
        self.assertTrue(any(run.bold and run.text == "bold" for run in intro.runs))
        self.assertEqual(d.core_properties.title, "Network Report")

    def test_each_numbered_list_restarts_at_one(self):
        import docx
        from docx.oxml.ns import qn
        r = self.save("docx", "# T\n1. a\n2. b\n\ntext between\n\n1. c\n2. d\n", "nums")
        d = docx.Document(str(r.path))
        ids = [p._p.pPr.numPr.numId.val for p in d.paragraphs if p.style.name == "List Number"]
        self.assertEqual(ids[0], ids[1])
        self.assertEqual(ids[2], ids[3])
        self.assertNotEqual(ids[1], ids[2])


@unittest.skipUnless(HAVE["pptx"], "python-pptx not installed")
class PptxTests(TmpCase):
    def titles(self, path):
        import pptx
        return [s.shapes.title.text_frame.text for s in pptx.Presentation(str(path)).slides]

    def test_pptx_slide_count_and_titles(self):
        r = self.save("pptx", DECK, "deck", slides=4)
        self.assertEqual(r.details["slides"], 4)
        self.assertEqual(self.titles(r.path), ["Marine Pollution", "Causes", "Steps", "Data"])
        self.assertEqual(r.warnings, [])  # the requested count was met, so no complaint

    def test_pptx_body_and_table_text_is_really_there(self):
        import pptx
        r = self.save("pptx", DECK, "deck")
        prs = pptx.Presentation(str(r.path))
        texts = []
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    texts.append(sh.text_frame.text)
                if getattr(sh, "has_table", False):
                    texts.extend(c.text for row in sh.table.rows for c in row.cells)
        blob = "\n".join(texts)
        for needle in ("Protecting our oceans", "Plastic waste", "Oil spills", "Reduce", "Tonnes", "2021"):
            self.assertIn(needle, blob)  # a table slide used to lose its table silently

    def test_slide_is_16_9_and_title_slide_is_styled(self):
        import pptx
        r = self.save("pptx", DECK, "deck")
        prs = pptx.Presentation(str(r.path))
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=2)

    def test_long_lists_are_split_over_slides(self):
        md = "# Deck\n## Big list\n" + "\n".join("- item number %d with some words" % i for i in range(20))
        r = self.save("pptx", md, "big")
        self.assertGreaterEqual(r.details["slides"], 3)
        self.assertIn("(cont.)", self.titles(r.path)[-1])

    def test_slide_count_mismatch_is_reported(self):
        r = self.save("pptx", DECK, "deck", slides=7)
        self.assertTrue(any("asked for 7 slides" in w for w in r.warnings))

    def test_content_without_a_title_still_gets_a_title_slide(self):
        r = self.save("pptx", "## One\n- a\n## Two\n- b", "x", title="My Talk")
        self.assertEqual(self.titles(r.path), ["My Talk", "One", "Two"])


@unittest.skipUnless(HAVE["xlsx"], "openpyxl not installed")
class XlsxTests(TmpCase):
    def test_workbook_sheets_values_and_formatting(self):
        import openpyxl
        r = self.save("xlsx", SHEET, "marks", allow_formulas=True)
        wb = openpyxl.load_workbook(str(r.path))
        self.assertEqual(wb.sheetnames, ["Students", "Notes"])
        ws = wb["Students"]
        self.assertEqual([c.value for c in ws[1]], ["Name", "Maths", "Science", "Total"])
        self.assertEqual([c.value for c in ws[2]], ["Asha", 90, 85, "=SUM(B2:C2)"])
        self.assertEqual(ws["D3"].value, "=SUM(B3:C3)")
        self.assertEqual(ws["B2"].data_type, "n")
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:D3")
        self.assertGreater(ws.column_dimensions["A"].width, 9)
        self.assertEqual(wb["Notes"]["B2"].value, 40)

    def test_formulas_only_when_requested(self):
        import openpyxl
        r = self.save("xlsx", SHEET, "nof", allow_formulas=False)
        ws = openpyxl.load_workbook(str(r.path))["Students"]
        self.assertEqual(ws["D2"].data_type, "s")
        self.assertEqual(ws["D2"].value, "=SUM(B{row}:C{row})")

    def test_spreadsheet_injection_stays_text(self):
        import openpyxl
        evil = ['=HYPERLINK("http://evil.example","x")', "=cmd|' /C calc'!A0", "=1+1", "@SUM(1)",
                "=WEBSERVICE(\"http://e\")", "=Sheet2!A1"]
        md = "# T\n| Name |\n|---|\n" + "\n".join("| %s |" % e.replace("|", "\\|") for e in evil)  # \| = literal pipe
        r = self.save("xlsx", md, "evil", allow_formulas=True)
        ws = openpyxl.load_workbook(str(r.path)).active
        for i, e in enumerate(evil, start=2):
            cell = ws.cell(row=i, column=1)
            if e == "=1+1":  # harmless arithmetic is allowed when formulas were requested
                self.assertEqual(cell.data_type, "f")
            else:
                self.assertNotEqual(cell.data_type, "f", e)
                self.assertEqual(cell.value, e)
        self.assertTrue(r.warnings)

    def test_numbers_text_and_leading_zeros(self):
        import openpyxl
        md = "# T\n| a | b | c | d | e |\n|---|---|---|---|---|\n| 007 | 1,200 | 85% | 3.5 | +91 98765 |"
        r = self.save("xlsx", md, "types")
        ws = openpyxl.load_workbook(str(r.path)).active
        self.assertEqual(ws["A2"].value, "007")
        self.assertEqual(ws["B2"].value, 1200)
        self.assertAlmostEqual(ws["C2"].value, 0.85)
        self.assertEqual(ws["C2"].number_format, "0%")
        self.assertEqual(ws["D2"].value, 3.5)
        self.assertEqual(ws["E2"].value, "+91 98765")

    def test_sheet_names_are_sanitised_and_unique(self):
        import openpyxl
        md = "## A/B:C?\n| x |\n|---|\n| 1 |\n## A/B:C?\n| x |\n|---|\n| 2 |\n## " + "L" * 50 + "\n| x |\n|---|\n| 3 |"
        r = self.save("xlsx", md, "names")
        names = openpyxl.load_workbook(str(r.path)).sheetnames
        self.assertEqual(len(set(n.lower() for n in names)), 3)
        self.assertTrue(all(len(n) <= 31 and not set(n) & set("[]:*?/\\") for n in names))

    def test_csv_answer_from_the_model_is_accepted(self):
        import openpyxl
        r = self.save("xlsx", "Name,Marks\nAsha,90\nRavi,78", "csv")
        ws = openpyxl.load_workbook(str(r.path)).active
        self.assertEqual([[c.value for c in row] for row in ws.iter_rows()], [["Name", "Marks"], ["Asha", 90], ["Ravi", 78]])

    def test_text_without_a_table_is_rejected(self):
        with self.assertRaises(D.DocError) as cm:
            self.save("xlsx", "just some words, no table", "bad")
        self.assertIn("table", str(cm.exception))

    def test_safe_formula_whitelist(self):
        for ok in ("=SUM(A1:A3)", "=B2*C2", "=ROUND(AVERAGE(B2:D2),1)", '=IF(B2>40,"pass","fail")'):
            self.assertTrue(D.safe_formula(ok), ok)
        for bad in ("=HYPERLINK(\"u\")", "=INDIRECT(A1)", "=Sheet1!A1", "=[1]a!A1", "=cmd|x", "=A1+FOO(1)",
                    "=SUM(A1:A3", "=A1;B1", "=" + "A1+" * 200 + "1", "SUM(A1)"):
            self.assertFalse(D.safe_formula(bad), bad)


# ═════════════════════════ file names, folders, failures ═════════════════════════
@unittest.skipUnless(ALL_LIBS, "document libraries not installed")
class SafetyTests(TmpCase):
    def test_filenames_with_spaces_and_extensions(self):
        r = self.save("pdf", REPORT, "Payment QR Report")
        self.assertEqual(r.path.name, "Payment QR Report.pdf")
        self.assertTrue(r.path.is_file())
        self.assertEqual(D.safe_filename("my file.PDF", "pdf"), "my file.pdf")
        self.assertEqual(D.safe_filename('we<ir>d:"name"?.docx', "docx"), "we_ir_d__name__.docx")

    def test_unsafe_filenames_are_rejected(self):
        for bad in ("../evil", "a/b", "a\\b", "..", ".hidden", "", "   ", "x.docx", "/etc/passwd"):
            with self.assertRaises(D.DocError, msg=bad):
                D.safe_filename(bad, "pdf")
        with self.assertRaises(D.DocError):
            self.save("pdf", REPORT, "../escape")
        self.assertEqual([p for p in self.tmp.rglob("escape*")], [])

    def test_existing_files_are_never_overwritten(self):
        a = self.save("docx", REPORT, "same")
        keep = a.path.read_bytes()
        b = self.save("docx", "# Other\ntext", "same")
        self.assertEqual(b.path.name, "same (1).docx")
        self.assertEqual(a.path.read_bytes(), keep)
        self.assertTrue(any("already existed" in w for w in b.warnings))
        with self.assertRaises(D.DocError):
            self.save("docx", REPORT, "same", on_exists="error")
        c = self.save("docx", "# Third\ntext", "same", on_exists="overwrite")
        self.assertEqual(c.path.name, "same.docx")

    def test_symlink_target_is_not_followed(self):
        outside = self.tmp / "outside.pdf"
        outside.write_bytes(b"precious")
        (self.work / "link.pdf").symlink_to(outside)
        with self.assertRaises(D.DocError):
            self.save("pdf", REPORT, "link", on_exists="overwrite")
        self.assertEqual(outside.read_bytes(), b"precious")

    def test_folders_outside_the_allowed_places_are_refused_with_advice(self):
        for target in ("/etc", "/root/.ssh", "/proc", str(self.tmp / "elsewhere")):
            with self.assertRaises(D.DocError) as cm:
                self.paths.resolve(("path", target))
            self.assertIn("outside", str(cm.exception))
            self.assertIn("AGENT WORK", str(cm.exception))  # suggests an accessible destination

    def test_traversal_in_the_folder_is_refused(self):
        with self.assertRaises(D.DocError):
            self.paths.resolve(("path", "../../etc"))
        with self.assertRaises(D.DocError):
            self.paths.resolve(("path", str(self.shared / ".." / ".." / "etc")))

    def test_symlink_inside_workspace_cannot_escape(self):
        (self.work / "sneaky").symlink_to(self.tmp, target_is_directory=True)
        with self.assertRaises(D.DocError):
            self.paths.resolve(("path", "sneaky"))

    def test_hidden_folders_and_non_folders_are_refused(self):
        with self.assertRaises(D.DocError):
            self.paths.resolve(("path", ".secret"))
        (self.work / "afile").write_text("x")
        with self.assertRaises(D.DocError) as cm:
            self.paths.resolve(("path", "afile/inside"))
        self.assertIn("not a folder", str(cm.exception))

    def test_relative_folder_is_created_inside_the_workspace(self):
        folder, _ = self.paths.resolve(("path", "out/reports"))
        r = D.save_document("pdf", REPORT, "r", folder, self.paths)
        self.assertEqual(r.path, self.work / "out" / "reports" / "r.pdf")

    def test_agent_work_folder_name_with_space(self):
        folder, _ = self.paths.resolve(("agent_work", None))
        self.assertEqual(folder, self.shared / "AGENT WORK")
        r = D.save_document("pdf", REPORT, "Y2B_Test", folder, self.paths)
        self.assertEqual(r.path, self.shared / "AGENT WORK" / "Y2B_Test.pdf")
        folder2, _ = self.paths.resolve(("path", str(self.shared / "AGENT WORK")))
        self.assertEqual(folder2, folder)

    def test_default_folder_prefers_agent_work_then_workspace(self):
        folder, note = self.paths.default_dir()
        self.assertEqual(folder, self.shared / "AGENT WORK")
        D.SHARED_ROOTS[:] = [str(self.tmp / "no" / "such" / "root")]
        folder, note = D.DocPaths(self.work).default_dir()
        self.assertEqual(folder, self.work / "documents")

    def test_docs_dir_option_is_allowed_and_is_the_default(self):
        extra = self.tmp / "my docs"
        p = D.DocPaths(self.work, extra)
        folder, _ = p.resolve(None)
        self.assertEqual(folder, extra.resolve())
        self.assertTrue(p.is_trusted(extra))

    def test_missing_android_storage_permission_explains_termux_setup_storage(self):
        D.SHARED_ROOTS[:] = [str(self.tmp / "storage" / "emulated" / "gone")]  # not mounted / not granted
        p = D.DocPaths(self.work)
        with self.assertRaises(D.StorageError) as cm:
            p.resolve(("agent_work", None))
        self.assertIn("termux-setup-storage", str(cm.exception))
        self.assertIn("workspace", str(cm.exception))
        with self.assertRaises(D.StorageError):
            p.resolve(("path", str(self.tmp / "storage" / "emulated" / "gone" / "AGENT WORK")))
        self.assertFalse((self.tmp / "storage" / "emulated" / "gone").exists())  # nothing was created

    def test_os_errors_become_useful_messages(self):
        for exc, expect in ((PermissionError(13, "denied"), "No permission"), (OSError(28, "full"), "out of free storage"),
                            (OSError(30, "ro"), "read-only")):
            with mock.patch("tempfile.mkstemp", side_effect=exc):
                with self.assertRaises(D.DocError) as cm:
                    self.save("pdf", REPORT, "x")
            self.assertIn(expect, str(cm.exception))
        self.assertEqual(os.listdir(str(self.work)), [])

    def test_failed_build_leaves_nothing_behind_and_never_succeeds(self):
        with mock.patch.dict(D._BUILDERS, {"pdf": mock.Mock(side_effect=RuntimeError("boom"))}):
            with self.assertRaises(D.DocError) as cm:
                self.save("pdf", REPORT, "x")
        self.assertIn("boom", str(cm.exception))
        self.assertEqual(os.listdir(str(self.work)), [])  # no half-written file, no temp file

    def test_empty_or_corrupt_output_is_rejected(self):
        def empty(path, blocks, **kw):
            open(path, "wb").close()

        def junk(path, blocks, **kw):
            with open(path, "wb") as f:
                f.write(b"not a docx at all")
        for fn, word in ((empty, "empty"), (junk, "reopened")):
            with mock.patch.dict(D._BUILDERS, {"docx": fn}):
                with self.assertRaises(D.DocError) as cm:
                    self.save("docx", REPORT, "x")
            self.assertIn(word, str(cm.exception))
        self.assertEqual(os.listdir(str(self.work)), [])

    def test_missing_library_gives_install_instructions(self):
        real = importlib.util.find_spec
        with mock.patch("importlib.util.find_spec", lambda n, *a: None if n == "reportlab" else real(n, *a)):
            self.assertEqual(D.missing_libraries(["pdf", "docx"]), ["reportlab"])
            with self.assertRaises(D.MissingLibraryError) as cm:
                self.save("pdf", REPORT, "x")
        self.assertIn("pip install reportlab", str(cm.exception))
        self.assertTrue(self.save("docx", REPORT, "still works").path.is_file())  # other formats are unaffected

    def test_oversized_content_is_rejected(self):
        with self.assertRaises(D.DocError):
            self.save("pdf", "x" * (D.MAX_MARKDOWN_CHARS + 1), "big")


# ═════════════════════════ understanding requests ═════════════════════════
class DetectionTests(unittest.TestCase):
    def test_the_example_requests(self):
        d = D.detect_request
        r = d("Create a PDF explaining Wireshark and save it in AGENT WORK.")
        self.assertEqual((r.formats, r.dest, r.topic), (["pdf"], ("agent_work", None), "Wireshark"))
        r = d("Create an editable Word document titled Payment QR with the content I provide.")
        self.assertEqual((r.formats, r.title, r.needs_content), (["docx"], "Payment QR", True))
        r = d("Make a 5-slide PowerPoint presentation about marine pollution.")
        self.assertEqual((r.formats, r.slides, r.topic), (["pptx"], 5, "marine pollution"))
        r = d("Create an Excel spreadsheet containing student names, marks, and total scores.")
        self.assertEqual((r.formats, r.allow_formulas), (["xlsx"], True))
        r = d("Create a professional project report in both PDF and Word formats.")
        self.assertEqual(r.formats, ["pdf", "docx"])
        r = d("Save a PDF called Y2B_Test.pdf in AGENT WORK.")
        self.assertEqual((r.formats, r.filename, r.dest, r.title), (["pdf"], "Y2B_Test.pdf", ("agent_work", None), None))

    def test_destinations(self):
        d = D.detect_request
        self.assertEqual(d("save a pdf called Notes.pdf to /storage/emulated/0/AGENT WORK/").dest,
                         ("path", "/storage/emulated/0/AGENT WORK/"))
        self.assertEqual(d("create /storage/emulated/0/AGENT WORK/Report.docx about tides").filename, "Report.docx")
        self.assertEqual(d("make a pdf about tides in ~/y2b_workspace/out please").dest, ("path", "~/y2b_workspace/out"))
        self.assertEqual(d("make a pdf about tides in the workspace").dest, ("workspace", None))
        self.assertIsNone(d("make a pdf about tides").dest)

    def test_slide_words_and_numbers(self):
        self.assertEqual(D.detect_request("make a five slide powerpoint on tides").slides, 5)
        self.assertEqual(D.detect_request("make a 12 slides deck about tides").slides, 12)
        self.assertIsNone(D.detect_request("make a powerpoint about tides").slides)

    def test_inline_content(self):
        r = D.detect_request("create a word document with the following content: Hello team, the meeting is at 5pm.")
        self.assertEqual(r.provided, "Hello team, the meeting is at 5pm.")

    def test_not_document_requests(self):
        for t in ("hi", "create hello.py that prints hello and run it", "write a python script that makes a pdf",
                  "how do I make a pdf", "what is an excel spreadsheet", "explain powerpoint to me",
                  "run sl", "create a web birthday calculator", "convert report.pdf to text", "write a word list app",
                  "make a game", "read notes.docx"):
            self.assertIsNone(D.detect_request(t), t)

    def test_file_names_follow_title_topic_or_explicit_name(self):
        r = D.detect_request("make a pdf about marine pollution")
        self.assertEqual(D.choose_filename(r, "pdf"), "marine_pollution.pdf")
        r = D.detect_request("make a word doc titled Payment QR")
        self.assertEqual(D.choose_filename(r, "docx"), "Payment_QR.docx")
        r = D.detect_request("save a pdf called My Notes.pdf")
        self.assertEqual(D.choose_filename(r, "pdf"), "My Notes.pdf")
        r = D.detect_request("create report.pdf and report.docx about tides")
        self.assertEqual(D.choose_filename(r, "docx"), "report.docx")


# ═════════════════════════ agent integration ═════════════════════════
@unittest.skipUnless(ALL_LIBS, "document libraries not installed")
class AgentDocTests(TmpCase):
    @classmethod
    def setUpClass(cls):
        cls.srv = fake_server.serve(0)
        cls.port = cls.srv.server_address[1]
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()
        cls.srv.server_close()

    def make_agent(self, auto=True):
        be = y2b.OpenAIBackend("http://127.0.0.1:%d" % self.port)
        be.start()
        llm = y2b.LLM(be, 4096, 256, 0.2)
        tools = y2b.Tools(str(self.work), auto=auto)
        return y2b.Agent(llm, tools, {"facts": []}, steps=4, fix=3)

    def handle(self, agent, text, answers=None):
        buf = io.StringIO()
        replies = iter(answers or [])

        def fake_input(prompt=""):
            print(prompt, end="")  # a real input() shows its prompt on stdout
            try:
                return next(replies)
            except StopIteration:
                raise EOFError
        with redirect_stdout(buf), mock.patch.object(builtins, "input", fake_input):
            agent.handle(text)
        return buf.getvalue()

    def test_pdf_saved_to_agent_work_and_path_reported(self):
        a = self.make_agent()
        out = self.handle(a, "Create a PDF explaining Wireshark and save it in AGENT WORK.")
        target = self.shared / "AGENT WORK" / "Wireshark.pdf"
        self.assertTrue(target.is_file())
        self.assertGreater(target.stat().st_size, 0)
        self.assertIn(str(target), out)
        self.assertIn("Saved PDF", out)
        text = pdf_text(target)
        if text is not None:
            self.assertIn("Wireshark Overview", text)
            self.assertIn("Capture packets", text)

    def test_explicit_android_path_with_a_space(self):
        a = self.make_agent()
        out = self.handle(a, "Save a PDF called Y2B_Test.pdf in %s" % (self.shared / "AGENT WORK"))
        self.assertTrue((self.shared / "AGENT WORK" / "Y2B_Test.pdf").is_file(), out)

    def test_word_with_content_the_user_provides(self):
        import docx
        a = self.make_agent()
        out = self.handle(a, "Create an editable Word document titled Payment QR with the content I provide.",
                          answers=["Pay to: Mahzend", "UPI: example@upi", ":wq"])
        files = list(self.shared.glob("**/*.docx")) + list(self.work.glob("**/*.docx"))
        self.assertEqual(len(files), 1, out)
        self.assertEqual(files[0].name, "Payment_QR.docx")
        d = docx.Document(str(files[0]))
        txt = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("Payment QR", txt)
        self.assertIn("UPI: example@upi", txt)

    def test_powerpoint_with_requested_slide_count(self):
        import pptx
        a = self.make_agent()
        out = self.handle(a, "Make a 5-slide PowerPoint presentation about marine pollution.")
        files = list(self.shared.glob("**/*.pptx"))
        self.assertEqual(len(files), 1, out)
        prs = pptx.Presentation(str(files[0]))
        self.assertEqual(len(prs.slides), 5)
        self.assertEqual(prs.slides[0].shapes.title.text_frame.text, "Marine Pollution")
        self.assertIn("5 slides", out)

    def test_excel_with_formulas_when_totals_are_asked_for(self):
        import openpyxl
        a = self.make_agent()
        out = self.handle(a, "Create an Excel spreadsheet containing student names, marks, and total scores.")
        files = list(self.shared.glob("**/*.xlsx"))
        self.assertEqual(len(files), 1, out)
        ws = openpyxl.load_workbook(str(files[0])).active
        self.assertEqual([c.value for c in ws[1]], ["Name", "Maths", "Science", "Total"])
        self.assertEqual(ws["D2"].value, "=SUM(B2:C2)")

    def test_both_pdf_and_word(self):
        a = self.make_agent()
        out = self.handle(a, "Create a professional project report in both PDF and Word formats.")
        names = sorted(p.suffix for p in self.shared.glob("AGENT WORK/*"))
        self.assertEqual(names, [".docx", ".pdf"], out)
        self.assertEqual(out.count("✓ saved"), 2)

    def test_default_destination_is_agent_work_when_available(self):
        a = self.make_agent()
        self.handle(a, "make a pdf about tides")
        self.assertTrue(list((self.shared / "AGENT WORK").glob("*.pdf")))

    def test_default_destination_falls_back_to_workspace_when_storage_is_unavailable(self):
        D.SHARED_ROOTS[:] = [str(self.tmp / "nothing")]
        a = self.make_agent()
        out = self.handle(a, "make a pdf about tides")
        self.assertTrue(list((self.work / "documents").glob("*.pdf")), out)

    def test_storage_permission_problem_is_reported_not_hidden(self):
        D.SHARED_ROOTS[:] = [str(self.tmp / "nothing")]
        a = self.make_agent()
        out = self.handle(a, "Create a PDF explaining Wireshark and save it in AGENT WORK.")
        self.assertIn("termux-setup-storage", out)
        self.assertNotIn("Saved", out)
        self.assertEqual(list(self.tmp.rglob("*.pdf")), [])

    def test_model_returning_nothing_is_not_reported_as_success(self):
        a = self.make_agent()
        out = self.handle(a, "make a pdf about empty-doc")
        self.assertIn("Nothing was saved", out)
        self.assertNotIn("✓ saved", out)
        self.assertEqual(list(self.tmp.rglob("*.pdf")), [])

    def test_spreadsheet_without_a_table_is_not_reported_as_success(self):
        a = self.make_agent()
        out = self.handle(a, "make an excel spreadsheet about prose-sheet")
        self.assertIn("Nothing was saved", out)
        self.assertIn("table", out)
        self.assertEqual(list(self.tmp.rglob("*.xlsx")), [])

    def test_a_build_failure_is_not_reported_as_success(self):
        a = self.make_agent()
        with mock.patch.dict(D._BUILDERS, {"pdf": mock.Mock(side_effect=RuntimeError("boom"))}):
            out = self.handle(a, "make a pdf about tides")
        self.assertIn("Nothing was saved", out)
        self.assertIn("boom", out)
        self.assertNotIn("Saved", out)
        self.assertEqual(list(self.tmp.rglob("*.pdf")), [])

    def test_post_save_check_catches_a_missing_file(self):
        a = self.make_agent()
        real = D.save_document

        def liar(*args, **kw):
            res = real(*args, **kw)
            os.unlink(str(res.path))  # pretend the file vanished after saving
            return res
        with mock.patch.object(D, "save_document", liar):
            out = self.handle(a, "make a pdf about tides")
        self.assertIn("Nothing was saved", out)
        self.assertNotIn("Saved PDF", out)

    def test_missing_library_message_comes_from_the_agent(self):
        a = self.make_agent()
        real = importlib.util.find_spec
        with mock.patch("importlib.util.find_spec", lambda n, *x: None if n == "pptx" else real(n, *x)):
            out = self.handle(a, "make a 5-slide PowerPoint about tides")
        self.assertIn("pip install python-pptx", out)
        self.assertEqual(list(self.tmp.rglob("*.pptx")), [])

    def test_confirmation_is_respected_when_not_in_auto_mode(self):
        a = self.make_agent(auto=False)
        out = self.handle(a, "make a pdf about tides", answers=["n"])
        self.assertIn("Nothing was saved", out)
        self.assertEqual(list(self.tmp.rglob("*.pdf")), [])
        out = self.handle(a, "make a pdf about tides", answers=["y"])
        self.assertEqual(len(list((self.shared / "AGENT WORK").glob("*.pdf"))), 1, out)

    def test_auto_mode_still_asks_for_untrusted_shared_storage_folders(self):
        a = self.make_agent(auto=True)
        (self.shared / "Download").mkdir()
        target = str(self.shared / "Download")
        out = self.handle(a, "make a pdf about tides in %s" % target, answers=["n"])
        self.assertEqual(list((self.shared / "Download").glob("*.pdf")), [], out)
        out = self.handle(a, "make a pdf about tides in %s" % target, answers=["y"])
        self.assertEqual(len(list((self.shared / "Download").glob("*.pdf"))), 1, out)

    def test_forbidden_folder_gives_an_error_and_a_suggestion(self):
        a = self.make_agent()
        out = self.handle(a, "make a pdf about tides in /etc")
        self.assertIn("outside", out)
        self.assertIn("AGENT WORK", out)
        self.assertFalse(os.path.exists("/etc/tides.pdf"))

    def test_existing_file_is_kept_and_new_one_gets_a_number(self):
        a = self.make_agent()
        self.handle(a, "make a pdf about tides")
        out = self.handle(a, "make a pdf about tides")
        names = sorted(p.name for p in (self.shared / "AGENT WORK").glob("*.pdf"))
        self.assertEqual(names, ["tides (1).pdf", "tides.pdf"], out)

    def test_it_asks_only_when_there_is_no_topic_at_all(self):
        a = self.make_agent()
        out = self.handle(a, "create a pdf", answers=["Wireshark basics"])
        self.assertIn("What should the PDF be about?", out)
        self.assertTrue((self.shared / "AGENT WORK" / "Wireshark_basics.pdf").is_file(), out)
        out = self.handle(a, "create a pdf", answers=[""])
        self.assertIn("nothing created", out)

    def test_existing_behaviour_is_untouched(self):
        a = self.make_agent()
        out = self.handle(a, "create calc.py with an addition function")
        self.assertTrue((self.work / "calc.py").exists(), out)
        self.assertEqual(list(self.tmp.rglob("*.pdf")), [])
        self.assertIn("Y2B", self.handle(a, "hi bro"))
        self.assertEqual(self.handle(a, "how do I make a pdf").count("✓ saved"), 0)

    def test_code_tool_still_cannot_write_outside_the_workspace(self):
        a = self.make_agent()
        r = a.tools.write(str(self.shared / "AGENT WORK" / "x.txt"), "hi")
        self.assertTrue(r.startswith("REFUSED"))
        self.assertFalse((self.shared / "AGENT WORK" / "x.txt").exists())

    def test_doctor_lists_document_status(self):
        class Args:
            work = str(self.work)
            docs_dir = None
        buf = io.StringIO()
        with redirect_stdout(buf):
            y2b.print_docs_status(Args)
        out = buf.getvalue()
        self.assertIn("reportlab", out)
        self.assertIn("AGENT WORK", out)


if __name__ == "__main__":
    unittest.main()
