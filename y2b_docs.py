#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Y2B Agent - offline document generation (PDF, Word, PowerPoint, Excel).

    Made by : Mahzend

Three layers, each usable on its own:

  1. parse_markdown(text)        -> blocks   a tiny Markdown dialect, parsed as DATA (nothing is executed)
  2. create_pdf / create_docx /
     create_pptx / create_xlsx   -> file     blocks -> a real document at a given path
  3. save_document(...)          -> DocResult  safe path checks, atomic write, reopen + verify

Everything works offline. The heavy libraries (reportlab, python-docx, python-pptx, openpyxl) are
imported lazily, so the rest of Y2B Agent keeps working (stdlib only) when they are not installed.
"""
from __future__ import annotations

import errno
import hashlib
import importlib.util
import os
import re
import tempfile
import zipfile
from pathlib import Path

AUTHOR = "Y2B Agent"
FORMATS = ("pdf", "docx", "pptx", "xlsx")
EXT = {"pdf": ".pdf", "docx": ".docx", "pptx": ".pptx", "xlsx": ".xlsx"}
LABEL = {"pdf": "PDF", "docx": "Word", "pptx": "PowerPoint", "xlsx": "Excel"}
LEGACY_EXT = {".doc": "docx", ".ppt": "pptx", ".xls": "xlsx"}
LIBS = {"pdf": ("reportlab", "reportlab"), "docx": ("docx", "python-docx"),
        "pptx": ("pptx", "python-pptx"), "xlsx": ("openpyxl", "openpyxl")}

AGENT_WORK_NAME = "AGENT WORK"
# Android shared storage as seen from Termux. Module-level so tests (and unusual setups) can change it.
SHARED_ROOTS = ["/storage/emulated/0", "/sdcard", "~/storage/shared"]
DEFAULT_SUBDIR = "documents"          # used inside the workspace when AGENT WORK is not available

MAX_MARKDOWN_CHARS = 400000
MAX_COLS = 40
MAX_ROWS = 20000
MAX_SLIDES = 60


class DocError(Exception):
    """Anything that should be shown to the person as a plain, useful message."""


class StorageError(DocError):
    pass


class MissingLibraryError(DocError):
    pass


# ═══════════════════════════ libraries ═══════════════════════════
def missing_libraries(formats=FORMATS):
    """pip names of the libraries needed for `formats` that are not installed."""
    out = []
    for f in formats:
        mod, pip = LIBS[f]
        if importlib.util.find_spec(mod) is None and pip not in out:
            out.append(pip)
    return out


def install_hint(missing):
    return ("Missing Python package%s: %s.\n"
            "Install %s (one-time, needs internet):  pip install %s\n"
            "On Termux, first run:  pkg install python-pip python-pillow python-lxml   (or just: bash install.sh)"
            % ("s" if len(missing) > 1 else "", ", ".join(missing),
               "them" if len(missing) > 1 else "it", " ".join(missing)))


def _need(fmt):
    mod, pip = LIBS[fmt]
    if importlib.util.find_spec(mod) is None:
        raise MissingLibraryError(install_hint([pip]))


# ═══════════════════════════ text helpers ═══════════════════════════
_CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def clean_text(s):
    s = (s or "").replace("\r\n", "\n").replace("\r", "\n")
    s = s.encode("utf-8", "ignore").decode("utf-8", "ignore")  # drops lone surrogates
    return _CTRL_RE.sub("", s)


def strip_fences(s):
    s = (s or "").strip("\n")
    m = re.match(r"^```[\w+.#-]*[ \t]*\n(.*?)\n?```\s*$", s, re.S)
    return m.group(1) if m else s


_INLINE_RE = re.compile(r"\*\*\*(?P<bi>.+?)\*\*\*|\*\*(?P<b>.+?)\*\*|\*(?!\s)(?P<i>[^*\n]+?)(?<!\s)\*|`(?P<c>[^`\n]+)`")


def parse_inline(text):
    """'a **b** *c*' -> [(text, bold, italic), ...]"""
    runs, pos = [], 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False, False))
        if m.group("bi") is not None:
            runs.append((m.group("bi"), True, True))
        elif m.group("b") is not None:
            runs.append((m.group("b"), True, False))
        elif m.group("i") is not None:
            runs.append((m.group("i"), False, True))
        else:
            runs.append((m.group("c"), False, False))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], False, False))
    return runs or [("", False, False)]


def plain(text):
    return "".join(r[0] for r in parse_inline(text))


# ═══════════════════════════ layer 1: markdown -> blocks ═══════════════════════════
# Block types:
#   {"type": "title", "text"}                {"type": "heading", "level": 1-4, "text"}
#   {"type": "paragraph", "text"}            {"type": "bullets"|"numbered", "items": [(text, level), ...]}
#   {"type": "table", "header": [...], "rows": [[...], ...]}      {"type": "pagebreak"}
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_BULLET_RE = re.compile(r"^(\s*)([-*+•●▪])\s+(.*)$")
_NUMBER_RE = re.compile(r"^(\s*)(\d{1,3})[.)]\s+(.*)$")
_HR_RE = re.compile(r"^(?:-{3,}|\*{3,}|_{3,}|={3,})$")
_PAGEBREAK_RE = re.compile(r"^(?:\\(?:pagebreak|newpage)|\[\s*page\s*-?\s*break\s*\]|<\s*page\s*-?\s*break\s*/?\s*>"
                           r"|<!--\s*page\s*-?\s*break\s*-->)$", re.I)
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(?:\|\s*:?-{2,}:?\s*)*\|?$")
_IMAGE_RE = re.compile(r"^!\[[^\]]*\]\([^)]*\)\s*$")


def _split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    return [c.replace("\\|", "|").strip() for c in re.split(r"(?<!\\)\|", s)]


def _is_structure(line, nxt=""):
    s = line.strip()
    return bool(not s or _HEADING_RE.match(s) or _HR_RE.match(s) or _PAGEBREAK_RE.match(s)
                or _BULLET_RE.match(line) or _NUMBER_RE.match(line) or s.startswith("|")
                or ("|" in s and _TABLE_SEP_RE.match(nxt.strip()) and "|" in nxt))


def parse_markdown(text, warnings=None):
    """Parse the Markdown dialect into blocks. Never raises on odd input; unknown syntax stays plain text."""
    warnings = warnings if warnings is not None else []
    lines = clean_text(strip_fences(text)).split("\n")
    n, i, blocks = len(lines), 0, []
    have_title = False
    dropped_images = 0

    while i < n:
        line = lines[i].rstrip()
        s = line.strip()
        nxt = lines[i + 1] if i + 1 < n else ""
        if not s or _HR_RE.match(s):
            i += 1
            continue
        if _PAGEBREAK_RE.match(s):
            blocks.append({"type": "pagebreak"})
            i += 1
            continue
        m = _HEADING_RE.match(s)
        if m:
            level, txt = len(m.group(1)), plain(m.group(2)).strip()
            if level == 1 and not have_title and not [b for b in blocks if b["type"] != "pagebreak"]:
                blocks.append({"type": "title", "text": txt})
                have_title = True
            else:
                blocks.append({"type": "heading", "level": min(level, 4), "text": txt})
            i += 1
            continue
        if s.startswith("|") or ("|" in s and _TABLE_SEP_RE.match(nxt.strip()) and "|" in nxt):
            rows = []
            while i < n and "|" in lines[i] and lines[i].strip():
                if not _TABLE_SEP_RE.match(lines[i].strip()):
                    rows.append(_split_row(lines[i]))
                i += 1
            if rows:
                width = min(MAX_COLS, max(len(r) for r in rows))
                rows = [(r + [""] * width)[:width] for r in rows]
                blocks.append({"type": "table", "header": rows[0], "rows": rows[1:]})
            continue
        mb, mn = _BULLET_RE.match(line), _NUMBER_RE.match(line)
        if mb or mn:
            kind = "bullets" if mb else "numbered"
            items, prev = [], -1
            while i < n:
                l = lines[i].rstrip()
                if not l.strip():
                    j = i + 1
                    while j < n and not lines[j].strip():
                        j += 1
                    if j < n and (_BULLET_RE if kind == "bullets" else _NUMBER_RE).match(lines[j]):
                        i = j
                        continue
                    break
                b2, n2 = _BULLET_RE.match(l), _NUMBER_RE.match(l)
                m2 = b2 if kind == "bullets" else n2
                if m2:
                    level = min(2, len(m2.group(1).expandtabs(4)) // 2, prev + 1)
                    items.append([m2.group(3).strip(), max(level, 0)])
                    prev = items[-1][1]
                    i += 1
                    continue
                if l[:1] in (" ", "\t") and items and not (b2 or n2) and not _is_structure(l.strip()):
                    items[-1][0] += " " + l.strip()
                    i += 1
                    continue
                break
            blocks.append({"type": kind, "items": [(t, lv) for t, lv in items]})
            continue
        para = []
        while i < n and not (para and _is_structure(lines[i], lines[i + 1] if i + 1 < n else "")):
            seg = lines[i].strip()
            if _IMAGE_RE.match(seg):
                dropped_images += 1
            else:
                para.append(re.sub(r"^>\s?", "", seg))
            i += 1
        if para:
            blocks.append({"type": "paragraph", "text": " ".join(para).strip()})
    if dropped_images:
        warnings.append("%d image(s) were left out (images are not supported yet)." % dropped_images)
    return blocks


def block_text(b):
    t = b["type"]
    if t in ("title", "heading", "paragraph"):
        return b["text"]
    if t in ("bullets", "numbered"):
        return "\n".join(x[0] for x in b["items"])
    if t == "table":
        return "\n".join(" ".join(r) for r in [b["header"]] + b["rows"])
    return ""


def all_text(blocks):
    return "\n".join(block_text(b) for b in blocks)


def validate_blocks(blocks):
    if not any(block_text(b).strip() for b in blocks):
        raise DocError("The document has no content. Give a topic or some text and try again.")
    for b in blocks:
        if b["type"] == "table" and len(b["rows"]) + 1 > MAX_ROWS:
            raise DocError("A table has more than %d rows; split it into several sheets/tables." % MAX_ROWS)


def csv_to_blocks(text):
    """Fallback for spreadsheets: a model sometimes answers with CSV instead of a table."""
    import csv
    lines = [l for l in clean_text(strip_fences(text)).split("\n") if l.strip()]
    for delim in (",", ";", "\t"):
        try:
            rows = list(csv.reader(lines, delimiter=delim))
        except csv.Error:
            continue
        if len(rows) >= 2 and len(rows[0]) >= 2 and sum(len(r) == len(rows[0]) for r in rows) >= max(2, int(len(rows) * .8)):
            w = min(MAX_COLS, len(rows[0]))
            rows = [([c.strip() for c in r] + [""] * w)[:w] for r in rows]
            return [{"type": "table", "header": rows[0], "rows": rows[1:]}]
    return []


def count_slides(blocks):
    return len(_blocks_to_slides(blocks))


def document_title(blocks, fallback=""):
    for b in blocks:
        if b["type"] == "title":
            return b["text"]
    for b in blocks:
        if b["type"] == "heading":
            return b["text"]
    return fallback


# ═══════════════════════════ file names and folders ═══════════════════════════
_BAD_NAME_CHARS = re.compile(r'[<>:"|?*\x00-\x1f]')


def safe_filename(name, fmt):
    """Validate a plain file name (no folders) and make sure it ends with the right extension."""
    name = (name or "").strip().strip("'\"`")
    if not name:
        raise DocError("The file name is empty.")
    if "/" in name or "\\" in name or ".." in name.replace("...", ""):
        raise DocError("File name '%s' must not contain folders or '..'. Choose the folder separately." % name)
    if name.startswith("."):
        raise DocError("File name '%s' must not start with a dot (hidden files are not allowed)." % name)
    low = name.lower()
    for other in FORMATS:
        if low.endswith(EXT[other]) and other != fmt:
            raise DocError("File name '%s' ends in %s but a %s file was requested." % (name, EXT[other], LABEL[fmt]))
    stem = name[:-len(EXT[fmt])] if low.endswith(EXT[fmt]) else name
    stem = re.sub(r"\s+", " ", _BAD_NAME_CHARS.sub("_", stem)).strip(" .")
    if not stem:
        raise DocError("The file name has no usable characters.")
    return stem[:100].rstrip(" .") + EXT[fmt]


def slugify(text, max_words=5):
    words = re.findall(r"[A-Za-z0-9À-ɏ]+", text or "")
    return "_".join(w[:24] for w in words[:max_words]) or "document"


def unique_path(path):
    """name.pdf -> name (1).pdf -> name (2).pdf, never an existing file."""
    p = Path(path)
    k = 1
    while True:
        cand = p.with_name("%s (%d)%s" % (p.stem, k, p.suffix))
        if not cand.exists():
            return cand
        k += 1


def _within(child, root):
    try:
        child.relative_to(root)
        return True
    except ValueError:
        return False


def _expand(p):
    return Path(os.path.expanduser(str(p)))


def is_termux():
    return "com.termux" in os.environ.get("PREFIX", "")


def storage_help(root=None):
    return ("Android storage is not accessible%s. In Termux run:  termux-setup-storage  "
            "(then tap Allow on the permission prompt), restart Y2B and try again. "
            "To save inside the workspace instead, say: save it in the workspace."
            % (" (%s)" % root if root else ""))


def _shared_roots():
    return [_expand(r) for r in SHARED_ROOTS]


def _root_usable(root):
    try:
        root = Path(root)
        if not root.is_dir() or not os.access(str(root), os.R_OK | os.W_OK | os.X_OK):
            return False
        os.listdir(str(root))
        return True
    except OSError:
        return False


def storage_problem(path):
    """If `path` is on Android shared storage that we cannot use, return help text (else None)."""
    p = _expand(path)
    try:
        rp = p.resolve()
    except OSError:
        rp = p
    for root in _shared_roots():
        try:
            rr = root.resolve()
        except OSError:
            rr = root
        if (_within(p, root) or _within(rp, rr)) and not _root_usable(root):
            return storage_help(str(root))
    return None


def agent_work_dir():
    """The AGENT WORK folder on the first usable Android shared-storage root (else the preferred path)."""
    cands = [r / AGENT_WORK_NAME for r in _shared_roots()]
    for c in cands:
        if _root_usable(c.parent):
            return c
    return cands[0]


def explain_os_error(exc, folder=None):
    """Turn an OSError into a message the person can act on."""
    folder = str(folder) if folder else "that folder"
    prob = storage_problem(folder) if folder != "that folder" else None
    code = getattr(exc, "errno", None)
    if isinstance(exc, PermissionError) or code in (errno.EACCES, errno.EPERM):
        return prob or ("No permission to write to %s. Pick a folder you own, e.g. the workspace "
                        "(say: save it in the workspace)." % folder)
    if code == errno.ENOSPC:
        return "The device is out of free storage, so the file could not be saved to %s." % folder
    if code == errno.EROFS:
        return "%s is read-only. Choose another folder." % folder
    if code == errno.ENAMETOOLONG:
        return "The file name is too long. Use a shorter name."
    if isinstance(exc, FileNotFoundError):
        return prob or "The folder %s does not exist and could not be created." % folder
    return prob or ("Could not write to %s: %s" % (folder, exc))


class DocPaths:
    """Where documents may be written. Allowed: the workspace, Android shared storage and --docs-dir."""

    def __init__(self, workspace, default_dir=None):
        self.workspace = Path(workspace).expanduser().resolve()
        self.override = Path(os.path.expanduser(str(default_dir))) if default_dir else None
        if self.override is not None and not self.override.is_absolute():
            self.override = self.workspace / self.override

    def roots(self):
        roots = [self.workspace] + _shared_roots()
        if self.override is not None:
            roots.append(self.override)
        out = []
        for r in roots:
            try:
                out.append(r.resolve())
            except OSError:
                out.append(r)
        return out

    def check_allowed(self, folder):
        try:
            rp = folder.resolve()
        except OSError:
            rp = folder
        for root in self.roots():
            if _within(rp, root):
                rel = rp.relative_to(root).parts
                if any(part.startswith(".") for part in rel):
                    raise DocError("Hidden folders (starting with a dot) are not allowed: %s" % folder)
                return rp
        raise DocError("%s is outside the folders Y2B may save documents to.\nAllowed: the workspace (%s), "
                       "Android shared storage (%s)%s.\nTry: save it in AGENT WORK, or in the workspace."
                       % (folder, self.workspace, SHARED_ROOTS[0],
                          ", and " + str(self.override) if self.override else ""))

    def check_writable(self, folder):
        """Validate (without creating anything) that `folder` can be used."""
        rp = self.check_allowed(folder)
        prob = storage_problem(rp)
        if prob:
            raise StorageError(prob)
        anc = rp
        while not anc.exists() and anc != anc.parent:
            anc = anc.parent
        if not anc.is_dir():
            raise DocError("%s is not a folder." % anc)
        if not os.access(str(anc), os.W_OK | os.X_OK):
            raise DocError(explain_os_error(PermissionError(errno.EACCES, "denied"), rp))
        return rp

    def is_trusted(self, folder):
        """True for the workspace, AGENT WORK and --docs-dir. Other shared-storage folders need a confirmation
        even in --auto mode, so auto-approve never writes outside the places the person set up."""
        try:
            rp = Path(folder).resolve()
        except OSError:
            rp = Path(folder)
        roots = [self.workspace, agent_work_dir()] + ([self.override] if self.override is not None else [])
        for r in roots:
            try:
                r = r.resolve()
            except OSError:
                pass
            if _within(rp, r):
                return True
        return False

    def default_dir(self):
        """(folder, note). AGENT WORK when Android storage is usable, else <workspace>/documents."""
        if self.override is not None:
            return self.override, None
        aw = agent_work_dir()
        if _root_usable(aw.parent):
            return aw, None
        note = None
        if is_termux():
            note = "AGENT WORK is not reachable yet. " + storage_help()
        return self.workspace / DEFAULT_SUBDIR, note

    def resolve(self, dest=None):
        """dest: None | ("agent_work", None) | ("workspace", None) | ("path", "<text>") -> (folder, note)"""
        kind, val = dest if dest else (None, None)
        if kind == "agent_work":
            folder, note = agent_work_dir(), None
        elif kind == "workspace":
            folder, note = self.workspace, None
        elif kind == "path":
            raw = str(val).strip().strip("'\"`")
            if "\x00" in raw or not raw:
                raise DocError("Invalid folder path.")
            p = Path(os.path.expanduser(raw))
            folder, note = (p if p.is_absolute() else self.workspace / p), None
        else:
            folder, note = self.default_dir()
        return self.check_writable(Path(folder)), note


# ═══════════════════════════ understanding a request ═══════════════════════════
_WORD_NUMS = {w: i for i, w in enumerate(
    "zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen "
    "sixteen seventeen eighteen nineteen twenty".split())}
_FORMAT_RE = {
    "pdf": re.compile(r"\bpdf\b|\.pdf\b", re.I),
    "docx": re.compile(
        r"\.docx?\b|\bdocx\b|\bms[ -]?word\b|\bmicrosoft word\b|\bword\s+(?:documents?|docs?|files?|formats?|versions?)\b"
        r"|\b(?:in|as|into|to)\s+(?:an?\s+)?(?:editable\s+)?word\b(?!s)"
        r"|\b(?:pdf|excel|pptx?|powerpoint|both)\s*(?:,|/|&|and|or)\s*(?:an?\s+)?(?:editable\s+)?word\b(?!s)"
        r"|\bword\s*(?:,|/|&|and|or)\s*(?:pdf|excel|pptx?|powerpoint)\b|\beditable\s+word\b", re.I),
    "pptx": re.compile(r"\.pptx?\b|\bpptx?\b|\bpowerpoint\b|\bslides?\b|\b(?:slide|pitch)[ -]?deck\b|\bpresentations?\b", re.I),
    "xlsx": re.compile(r"\.xlsx?\b|\bxlsx?\b|\bexcel\b|\bspreadsheets?\b|\bworkbooks?\b", re.I),
}
_VERB_RE = re.compile(r"\b(create|make|write|build|generate|save|prepare|produce|export|draft|compile|give me|i need|i want|"
                      r"need an?|want an?)\b", re.I)
_CONVERT_RE = re.compile(r"\b(?:put|turn|convert|export|save)\b.{0,60}?\b(?:into|to|as|in)\s+(?:an?\s+|the\s+)?"
                         r"(?:pdf|word|docx|powerpoint|pptx|excel|xlsx|spreadsheet)\b", re.I)
_QUESTION_RE = re.compile(r"^\s*(how|what|why|when|where|who|which|explain|tell|describe)\b", re.I)
_CODE_NOUN_RE = re.compile(r"\b(script|program|code|function|module|bot|api|app|application|game|website|web ?page|"
                           r"python|javascript|bash|node)\b|\.(?:py|js|sh)\b", re.I)
_FILE_REF_RE = re.compile(r"\b(?:convert|open|read|parse|extract|print|view)\s+\S+\.(?:pdf|docx?|pptx?|xlsx?)\b", re.I)
_DOCFILE = r"(?:pdf|docx?|pptx?|xlsx?)"
_ABS_FILE_RE = re.compile(r"(?<![\w/])((?:/|~/)[^\n\"'`]*?)/([^/\n\"'`]+\.%s)\b" % _DOCFILE, re.I)
_QUOTED_FILE_RE = re.compile(r"[\"'`“‘]([^\"'`”’\n/\\]+?\.%s)[\"'`”’]" % _DOCFILE, re.I)
_NAMED_FILE_RE = re.compile(r"\b(?:called|named|name\s+it|file\s*name\s*[:=]?)\s+[\"'`\u201c]?"
                            r"([\w][^\n\"'`/\\\u201d]{0,80}?\.%s)\b" % _DOCFILE, re.I)
_PLAIN_FILE_RE = re.compile(r"(?<![\w/\\.~-])([\w][\w\-]*(?:\.[\w\-]+)*\.%s)\b" % _DOCFILE, re.I)
_AGENTWORK_RE = re.compile(r"\bagent[\s_\-]*work\b", re.I)
_WORKSPACE_RE = re.compile(r"\b(?:in|to|into|inside)\s+(?:the\s+|my\s+)?(?:current\s+folder|workspace|work\s*folder|work\s*space)\b", re.I)
_PATH_RE = re.compile(r"\b(?:in|to|into|at|under|inside)\s+(?:the\s+)?(?:folder\s+|directory\s+|dir\s+)?[\"'`]?"
                      r"((?:/|~/)[^\"'`\n]+?)[\"'`]?(?=\s+(?:and|then|with|as|please|called|named)\b|[,;]|[.!?]?\s*$)", re.I)
_TITLE_RE = re.compile(r"\b(?:titled|title\s*[:=]?|called|named)\s+[\"“'`]?(.+?)[\"”'`]?"
                       r"(?=$|[.,;!?]|\s+(?:with|and|containing|that|which|using|for|in|on|at|saved?|please|the\s+content|as)\b)", re.I)
_TOPIC_RE = re.compile(r"\b(?:about|on|explaining|explain|regarding|covering|describing|containing|of|for)\s+(.+)", re.I)
_TOPIC_STOP_RE = re.compile(r"\b(and|save|saved|in|into|to|with|using|as|that|which|then|please)\b|[,.;!?]", re.I)
_SLIDES_RE = re.compile(r"\b(\d{1,2}|%s)\s*[- ]?\s*slides?\b|\bslides?\s*(?:count)?\s*[:=]\s*(\d{1,2})\b" % "|".join(_WORD_NUMS), re.I)
_INLINE_CONTENT_RE = re.compile(r"\b(?:with|using)\s+(?:the\s+)?(?:following\s+)?(?:content|text)\s*[:\-]\s*(.+)$|"
                                r"\b(?:content|text)\s*[:\-]\s*\n(.+)$", re.I | re.S)
_NEEDS_CONTENT_RE = re.compile(
    r"\bcontent\s+(?:i|that i)(?:'ll|\s+will)?\s*(?:provide|give|paste|send|share|type)\b|\bmy\s+own\s+(?:content|text)\b"
    r"|\b(?:with|using)\s+(?:my|the\s+following|this|the\s+below)\s+(?:content|text)\b|\bcontent\s+below\b", re.I)
_FORMULA_RE = re.compile(r"\bformulas?\b|\btotals?\b|\bsum\b|\baverages?\b|\bcalculat\w*|\bsubtotals?\b", re.I)
_FILLER = set("a an the me my in on it to and or for of with that this please create make write build generate save "
              "prepare produce export draft compile new file named called want need can you i us then also pdf word docx "
              "powerpoint pptx excel xlsx spreadsheet presentation slides slide document documents editable both "
              "format formats agent work folder".split())


class DocRequest:
    def __init__(self):
        self.formats = []
        self.filename = None       # as typed by the person (may be None)
        self.title = None
        self.dest = None           # None | ("agent_work", None) | ("workspace", None) | ("path", text)
        self.slides = None
        self.provided = None       # content the person gave inline
        self.needs_content = False
        self.allow_formulas = False
        self.topic = ""

    def __repr__(self):
        return "DocRequest(%s)" % ", ".join("%s=%r" % kv for kv in sorted(vars(self).items()))


def _first_format_pos(text):
    pos = [m.start() for f in FORMATS for m in [_FORMAT_RE[f].search(text)] if m]
    return min(pos) if pos else -1


def detect_request(text):
    """Decide whether `text` asks for a PDF / Word / PowerPoint / Excel file. Returns DocRequest or None."""
    if not text or _QUESTION_RE.match(text):
        return None
    formats = [f for f in FORMATS if _FORMAT_RE[f].search(text)]
    if not formats:
        return None
    if not (_VERB_RE.search(text) or _CONVERT_RE.search(text)) or _FILE_REF_RE.search(text):
        return None
    head = text[:_first_format_pos(text)]
    if _CODE_NOUN_RE.search(head) or re.search(r"\b\w+\.(?:py|js|sh)\b", text):
        return None  # "write a python script that makes a pdf" is a coding request, not a document request

    r = DocRequest()
    r.formats = formats
    path_hint = None
    m = _ABS_FILE_RE.search(text)
    if m:
        path_hint, r.filename = m.group(1), m.group(2)
    else:
        m = _QUOTED_FILE_RE.search(text) or _NAMED_FILE_RE.search(text) or _PLAIN_FILE_RE.search(text)
        if m:
            r.filename = m.group(1)
    if r.filename:
        ext = os.path.splitext(r.filename)[1].lower()
        fmt = LEGACY_EXT.get(ext) or next((f for f in FORMATS if EXT[f] == ext), None)
        if fmt and fmt not in r.formats:
            r.formats.append(fmt)
        if fmt:
            r.formats = [f for f in r.formats if f == fmt or len(r.formats) > 1] if len(r.formats) == 1 else r.formats
    # destination
    pm = _PATH_RE.search(text)
    if path_hint:
        r.dest = ("path", path_hint)
    elif pm:
        r.dest = ("path", pm.group(1).strip())
    elif _AGENTWORK_RE.search(text):
        r.dest = ("agent_work", None)
    elif _WORKSPACE_RE.search(text):
        r.dest = ("workspace", None)
    # title
    m = _TITLE_RE.search(text)
    if m:
        t = m.group(1).strip(" .\"'")
        is_filename = re.match(r"\.%s\b" % _DOCFILE, text[m.end(1):], re.I)  # "called Notes.pdf" names a file
        if t and not is_filename and len(t) <= 120:
            r.title = t
    # slides
    m = _SLIDES_RE.search(text)
    if m and "pptx" in r.formats:
        v = (m.group(1) or m.group(2) or "").lower()
        n = int(v) if v.isdigit() else _WORD_NUMS.get(v)
        if n and 1 <= n <= MAX_SLIDES:
            r.slides = n
    # content supplied by the person
    m = _INLINE_CONTENT_RE.search(text)
    if m:
        body = (m.group(1) or m.group(2) or "").strip()
        if len(body) >= 15:
            r.provided = body
    r.needs_content = bool(_NEEDS_CONTENT_RE.search(text)) and not r.provided
    r.allow_formulas = "xlsx" in r.formats and bool(_FORMULA_RE.search(text))
    clean = text
    for rx in (_ABS_FILE_RE, _QUOTED_FILE_RE, _NAMED_FILE_RE, _PLAIN_FILE_RE, _PATH_RE, _AGENTWORK_RE):
        clean = rx.sub(" ", clean)
    r.topic = r.title or _topic(clean)
    return r


def _topic(text):
    m = _TOPIC_RE.search(text)
    chunk = m.group(1) if m else text
    chunk = _TOPIC_STOP_RE.split(chunk)[0] if m else chunk
    words = [w for w in re.findall(r"[A-Za-z0-9À-ɏ]+", chunk) if w.lower() not in _FILLER]
    if not words and m:
        words = [w for w in re.findall(r"[A-Za-z0-9À-ɏ]+", text) if w.lower() not in _FILLER]
    return " ".join(words[:5])


def choose_filename(req, fmt, title_hint=""):
    """The file name Y2B will use for `fmt` (explicit name > title > topic)."""
    if req.filename:
        stem = os.path.splitext(req.filename)[0]
        return safe_filename(stem + EXT[fmt], fmt)
    base = req.title or req.topic or title_hint or "document"
    return safe_filename(slugify(base) + EXT[fmt], fmt)


# ═══════════════════════════ prompts for the language model ═══════════════════════════
DOC_SYSTEM = ("You are a document generator. Reply with ONLY the document content in simple Markdown: no commentary, "
              "no introduction, no closing remarks, no fences.\n"
              "Syntax: '# Title' on the first line, '## Section', '### Subsection'. Bullets start with '- ', numbered "
              "items with '1. '. Tables use '| a | b |' rows with a '|---|---|' line under the header. "
              "Use **bold** sparingly. Put [pagebreak] on its own line for a page break. No images, no HTML.\n")


def generation_system(fmt, slides=None, allow_formulas=False):
    if fmt == "pptx":
        n = ("Make exactly %d slides in total, counting the title slide. " % slides) if slides else \
            "Make 5 to 8 slides in total. "
        return DOC_SYSTEM + ("This is a slide presentation. Slide 1 is the title slide: '# Presentation title' followed by one "
                             "short subtitle line. Every other slide starts with '## Slide title' followed by 3 to 5 short "
                             "bullet lines ('- '). " + n + "Keep every bullet under 15 words.")
    if fmt == "xlsx":
        s = DOC_SYSTEM + ("This is a spreadsheet. Start each worksheet with '## Sheet name' and then ONE table whose first row "
                          "is the header. One value per cell, plain numbers without units or currency signs. Give sample "
                          "rows if no data was supplied.")
        if allow_formulas:
            s += (" Calculated columns may use Excel formulas such as =SUM(B{row}:D{row}) or =AVERAGE(B{row}:D{row}); "
                  "write {row} literally and it becomes the row number. Only use SUM, AVERAGE, MIN, MAX, COUNT, ROUND, IF.")
        else:
            s += " Write calculated values as plain numbers, not formulas."
        return s
    return DOC_SYSTEM + ("Write a professional, well organised document: a title, a short introduction, clear sections with "
                         "headings, bullets or a table where useful, and a brief conclusion. About one to two pages.")


# ═══════════════════════════ layer 3: save + verify ═══════════════════════════
class DocResult:
    def __init__(self, path, fmt, size, details, warnings):
        self.path, self.fmt, self.size = Path(path), fmt, size
        self.details, self.warnings = details, warnings

    def summary(self):
        d = self.details
        bits = []
        if "pages" in d:
            bits.append("%d page%s" % (d["pages"], "" if d["pages"] == 1 else "s"))
        if "slides" in d:
            bits.append("%d slide%s" % (d["slides"], "" if d["slides"] == 1 else "s"))
        if "sheets" in d:
            bits.append("%d sheet%s" % (len(d["sheets"]), "" if len(d["sheets"]) == 1 else "s"))
        bits.append("%.1f KB" % (self.size / 1024.0))
        return ", ".join(bits)


def prepare_blocks(fmt, markdown, title=None, warnings=None):
    """Markdown text -> validated blocks (adds the title when the content has none)."""
    if len(markdown or "") > MAX_MARKDOWN_CHARS:
        raise DocError("The content is too large (%d characters, limit %d)." % (len(markdown), MAX_MARKDOWN_CHARS))
    blocks = parse_markdown(markdown, warnings)
    if fmt == "xlsx" and not any(b["type"] == "table" for b in blocks):
        blocks = csv_to_blocks(markdown) or blocks
        if not any(b["type"] == "table" for b in blocks):
            raise DocError("A spreadsheet needs a table (rows and columns), but the content has none.")
    validate_blocks(blocks)
    if not any(b["type"] == "title" for b in blocks):
        t = title or (document_title(blocks) if fmt == "pptx" else "")
        if t:
            blocks.insert(0, {"type": "title", "text": plain(t).strip()})
    return blocks


def plan_target(folder, filename):
    """(target path, exists) for a validated folder + file name."""
    target = Path(folder) / filename
    return target, target.exists() or target.is_symlink()


def verify_file(path, fmt):
    """Reopen the finished file with its own library. Raises DocError if it is missing, empty or unreadable."""
    p = Path(path)
    if not p.is_file():
        raise DocError("The file was not created: %s" % p)
    size = p.stat().st_size
    if size <= 0:
        raise DocError("The file is empty: %s" % p)
    details = {}
    try:
        if fmt == "pdf":
            with open(str(p), "rb") as f:
                head = f.read(8)
                f.seek(max(0, size - 1024))
                tail = f.read()
            if not head.startswith(b"%PDF-") or b"%%EOF" not in tail:
                raise ValueError("not a complete PDF")
        else:
            if not zipfile.is_zipfile(str(p)):
                raise ValueError("not a valid Office file")
            _need(fmt)
            if fmt == "docx":
                import docx
                d = docx.Document(str(p))
                details = {"paragraphs": len(d.paragraphs), "tables": len(d.tables)}
            elif fmt == "pptx":
                import pptx
                pr = pptx.Presentation(str(p))
                details = {"slides": len(pr.slides)}
            else:
                import openpyxl
                wb = openpyxl.load_workbook(str(p))
                details = {"sheets": list(wb.sheetnames)}
    except MissingLibraryError:
        raise
    except Exception as e:
        raise DocError("The saved file could not be reopened, so it was not kept (%s)." % e)
    details["size"] = size
    return details


_BUILDERS = {}   # filled in below: fmt -> create_* function


def save_document(fmt, markdown, filename, folder, paths, title=None, allow_formulas=False, slides=None,
                  on_exists="rename"):
    """Create the document and save it safely. Returns a DocResult only after the file was verified.

    folder      an already validated folder (see DocPaths.resolve)
    on_exists   "rename" (default, never overwrites), "overwrite" or "error"
    """
    if fmt not in FORMATS:
        raise DocError("Unsupported format '%s'. Supported: %s." % (fmt, ", ".join(FORMATS)))
    _need(fmt)
    folder = paths.check_writable(Path(folder))
    name = safe_filename(filename, fmt)
    warnings = []
    blocks = prepare_blocks(fmt, markdown, title, warnings)
    target, exists = plan_target(folder, name)
    if exists:
        if target.is_symlink() or not target.is_file():
            raise DocError("%s exists and is not a regular file, so it will not be replaced." % target)
        if on_exists == "error":
            raise DocError("%s already exists." % target)
        if on_exists == "rename":
            target = unique_path(target)
            warnings.append("%s already existed, so the new file was saved as %s." % (name, target.name))
    tmp = None
    try:
        try:
            folder.mkdir(parents=True, exist_ok=True)
            fd, tmp = tempfile.mkstemp(prefix=".y2b-", suffix=EXT[fmt], dir=str(folder))
            os.close(fd)
        except OSError as e:
            raise DocError(explain_os_error(e, folder))
        kw = {"slides_requested": slides} if fmt == "pptx" else {}
        extra = _BUILDERS[fmt](tmp, blocks, title=title or document_title(blocks), allow_formulas=allow_formulas,
                               warnings=warnings, **kw)
        details = verify_file(tmp, fmt)
        details.update({k: v for k, v in (extra or {}).items() if k not in details})
        try:
            os.replace(tmp, str(target))
            tmp = None
        except OSError as e:
            raise DocError(explain_os_error(e, folder))
        if not target.is_file() or target.stat().st_size <= 0:
            raise DocError("The file could not be confirmed after saving: %s" % target)
        return DocResult(target, fmt, target.stat().st_size, details, warnings)
    except DocError:
        raise
    except MemoryError:
        raise DocError("Not enough memory to build this document. Try shorter content.")
    except OSError as e:
        raise DocError(explain_os_error(e, folder))
    except Exception as e:  # a library failed: report it, never pretend it worked
        raise DocError("Could not create the %s file: %s: %s" % (LABEL[fmt], type(e).__name__, e))
    finally:
        if tmp:
            try:
                os.unlink(tmp)
            except OSError:
                pass


def docs_status(paths):
    """Lines for `y2b --doctor`."""
    lines = []
    for f in FORMATS:
        mod, pip = LIBS[f]
        ok = importlib.util.find_spec(mod) is not None
        ver = ""
        if ok:
            try:
                from importlib.metadata import version
                ver = " " + version(pip)
            except Exception:
                pass
        lines.append((LABEL[f], pip + ver if ok else None))
    folder, note = paths.default_dir()
    try:
        paths.check_writable(folder)
        state = "ready"
    except DocError as e:
        state = "not usable: " + str(e).split("\n")[0]
    return lines, folder, state, note


# ═══════════════════════════ layer 2: renderers ═══════════════════════════
NAVY = "1F3A5F"
ACCENT = "2E75B6"
LIGHT = "EAF1FB"

# ---------------------------------------------------------------- PDF
_PDF_FONT_SETS = [
    # (regular, bold, italic, bold-italic)
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"),
    ("/usr/share/fonts/TTF/DejaVuSans.ttf", "/usr/share/fonts/TTF/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/TTF/DejaVuSans-Oblique.ttf", "/usr/share/fonts/TTF/DejaVuSans-BoldOblique.ttf"),
    ("$PREFIX/share/fonts/TTF/DejaVuSans.ttf", "$PREFIX/share/fonts/TTF/DejaVuSans-Bold.ttf",
     "$PREFIX/share/fonts/TTF/DejaVuSans-Oblique.ttf", "$PREFIX/share/fonts/TTF/DejaVuSans-BoldOblique.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf"),
    ("/system/fonts/NotoSans-Regular.ttf", "/system/fonts/NotoSans-Bold.ttf",
     "/system/fonts/NotoSans-Italic.ttf", "/system/fonts/NotoSans-BoldItalic.ttf"),
    ("/system/fonts/Roboto-Regular.ttf", "/system/fonts/Roboto-Bold.ttf",
     "/system/fonts/Roboto-Italic.ttf", "/system/fonts/Roboto-BoldItalic.ttf"),
]
_FONT_CACHE = {}


def _pdf_font_candidates():
    env = os.environ.get("Y2B_PDF_FONT")  # path to a regular .ttf; Bold/Italic siblings are optional
    if env:
        base = os.path.expanduser(env)
        stem, ext = os.path.splitext(base)
        yield (base, stem + "-Bold" + ext, stem + "-Italic" + ext, stem + "-BoldItalic" + ext)
    for fs in _PDF_FONT_SETS:
        yield tuple(os.path.expandvars(p) for p in fs)


def pdf_font():
    """(family name, glyph-checker or None). A Unicode TrueType family when one is installed, else Helvetica."""
    if "font" in _FONT_CACHE:
        return _FONT_CACHE["font"]
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    result = ("Helvetica", None)
    for n, fs in enumerate(_pdf_font_candidates()):
        if not os.path.isfile(fs[0]):
            continue
        fam = "Y2BFont%d" % n
        try:
            regular = TTFont(fam, fs[0])
            pdfmetrics.registerFont(regular)
            names = [fam]
            for suffix, path in zip(("-B", "-I", "-BI"), fs[1:]):
                if os.path.isfile(path):
                    pdfmetrics.registerFont(TTFont(fam + suffix, path))
                    names.append(fam + suffix)
                else:
                    names.append(fam)  # missing style: fall back to the regular face
            pdfmetrics.registerFontFamily(fam, normal=names[0], bold=names[1] if len(names) > 1 else fam,
                                          italic=names[2] if len(names) > 2 else fam,
                                          boldItalic=names[3] if len(names) > 3 else fam)
            cmap = regular.face.charToGlyph
            result = (fam, lambda ch, cmap=cmap: ord(ch) in cmap)
            break
        except Exception:
            continue
    _FONT_CACHE["font"] = result
    return result


def _helvetica_ok(ch):
    try:
        ch.encode("cp1252")
        return True
    except UnicodeEncodeError:
        return False


def _pdf_sanitize(text, ok):
    """Replace characters the chosen font cannot draw (instead of printing black boxes). Returns (text, n_replaced)."""
    bad = 0
    out = []
    for ch in text:
        if ch in "\n\t" or ok(ch):
            out.append(ch)
        else:
            out.append("?")
            bad += 1
    return "".join(out), bad


def _rl_markup(text, ok):
    from xml.sax.saxutils import escape
    out, bad = [], 0
    for t, bold, ital in parse_inline(text):
        t, b = _pdf_sanitize(t, ok)
        bad += b
        t = escape(t)
        if bold:
            t = "<b>%s</b>" % t
        if ital:
            t = "<i>%s</i>" % t
        out.append(t)
    return "".join(out), bad


_CJK_RE = re.compile("[⺀-鿿가-힯＀-￯]")


def create_pdf(path, blocks, title="", allow_formulas=False, warnings=None):
    warnings = warnings if warnings is not None else []
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    fam, has = pdf_font()
    ok = has or _helvetica_ok
    bold = fam + "-B" if fam != "Helvetica" else "Helvetica-Bold"
    try:
        from reportlab.pdfbase import pdfmetrics
        pdfmetrics.getFont(bold)
    except Exception:
        bold = fam
    wrap = "CJK" if _CJK_RE.search(all_text(blocks)) else "LTR"
    shift = 1 if any(b["type"] == "title" for b in blocks) else 0  # "## Section" is top level under a "# Title"
    navy, accent = colors.HexColor("#" + NAVY), colors.HexColor("#" + ACCENT)

    def style(name, **kw):
        base = dict(fontName=fam, fontSize=10.5, leading=15, textColor=colors.HexColor("#222222"), alignment=TA_LEFT,
                    wordWrap=wrap, spaceAfter=6)
        base.update(kw)
        return ParagraphStyle(name, **base)

    st_title = style("Title", fontName=bold, fontSize=25, leading=30, textColor=navy, spaceAfter=8)
    st_body = style("Body")
    st_h = {1: style("H1", fontName=bold, fontSize=16, leading=20, textColor=navy, spaceBefore=12, spaceAfter=6, keepWithNext=1),
            2: style("H2", fontName=bold, fontSize=13, leading=17, textColor=accent, spaceBefore=9, spaceAfter=4, keepWithNext=1),
            3: style("H3", fontName=bold, fontSize=11.5, leading=15, textColor=navy, spaceBefore=6, spaceAfter=3, keepWithNext=1),
            4: style("H4", fontName=bold, fontSize=10.5, leading=14, textColor=navy, spaceBefore=4, spaceAfter=2, keepWithNext=1)}
    st_cell = style("Cell", fontSize=9, leading=12, spaceAfter=0)
    st_head = style("CellHead", fontName=bold, fontSize=9, leading=12, spaceAfter=0, textColor=colors.white)

    replaced = [0]

    def para(text, st, **kw):
        m, bad = _rl_markup(text, ok)
        replaced[0] += bad
        return Paragraph(m or "&nbsp;", st, **kw)

    story = []
    for b in blocks:
        t = b["type"]
        if t == "title":
            story.append(para(b["text"], st_title))
            story.append(HRFlowable(width="100%", thickness=1.4, color=accent, spaceBefore=0, spaceAfter=10))
        elif t == "heading":
            story.append(para(b["text"], st_h[max(1, b["level"] - shift)]))
        elif t == "paragraph":
            story.append(para(b["text"], st_body))
        elif t in ("bullets", "numbered"):
            counters = {}
            for text, level in b["items"]:
                counters[level] = counters.get(level, 0) + 1
                for deeper in [k for k in counters if k > level]:
                    del counters[deeper]
                mark = ("•" if level == 0 else "–") if t == "bullets" else "%d." % counters[level]
                indent = 0.9 * cm + level * 0.8 * cm
                st = style("li%s%d" % (t, level), leftIndent=indent, bulletIndent=indent - 0.55 * cm, spaceAfter=3)
                story.append(para(text, st, bulletText=mark))
            story.append(Spacer(1, 4))
        elif t == "table":
            width = A4[0] - 4 * cm
            ncol = len(b["header"])
            lens = [max([len(plain(r[c])) for r in [b["header"]] + b["rows"]] + [3]) for c in range(ncol)]
            weights = [min(max(l, 6), 40) for l in lens]
            cw = [width * w / float(sum(weights)) for w in weights]
            data = [[para(c, st_head) for c in b["header"]]] + [[para(c, st_cell) for c in r] for r in b["rows"]]
            tb = Table(data, colWidths=cw, repeatRows=1)
            ts = [("BACKGROUND", (0, 0), (-1, 0), navy), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C4D6")),
                  ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                  ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
            for r in range(2, len(data), 2):
                ts.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#" + LIGHT)))
            tb.setStyle(TableStyle(ts))
            story.append(tb)
            story.append(Spacer(1, 10))
        elif t == "pagebreak":
            story.append(PageBreak())
    while story and isinstance(story[-1], PageBreak):
        story.pop()
    if not story:
        raise DocError("The document has no content.")

    pages = [0]
    doc_title = plain(title or document_title(blocks, "Document"))

    def footer(canvas, doc):
        pages[0] = max(pages[0], doc.page)
        canvas.saveState()
        canvas.setFont(fam, 8)
        canvas.setFillColor(colors.HexColor("#6B7785"))
        canvas.drawCentredString(A4[0] / 2.0, 1.0 * cm, "%d" % doc.page)
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                            bottomMargin=2 * cm, title=_pdf_sanitize(doc_title, ok)[0], author=AUTHOR, creator=AUTHOR)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    if replaced[0]:
        warnings.append("%d character(s) could not be drawn with the available PDF font and were shown as '?'. "
                        "Install a Unicode font (see README, Known limitations) or use the Word format." % replaced[0])
    return {"pages": pages[0]}


# ---------------------------------------------------------------- Word
def _docx_runs(paragraph, text, size=None, color=None, force_bold=False):
    from docx.shared import Pt, RGBColor
    for t, bold, ital in parse_inline(text):
        run = paragraph.add_run(t)
        run.bold = True if (bold or force_bold) else None
        run.italic = True if ital else None
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def _docx_restart_numbering(d):
    """Return a function that gives each numbered list its own numbering instance (so every list starts at 1)."""
    def new_list():
        numbering = d.part.numbering_part.numbering_definitions._numbering
        style = d.styles["List Number"]
        num_id = style.element.pPr.numPr.numId.val
        abstract = numbering.num_having_numId(num_id).abstractNumId.val
        num = numbering.add_num(abstract)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId
    return new_list


def create_docx(path, blocks, title="", allow_formulas=False, warnings=None):
    warnings = warnings if warnings is not None else []
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.3)
        s.top_margin = s.bottom_margin = Cm(2.2)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Title", 26, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 14, ACCENT),
                              ("Heading 3", 12, NAVY), ("Heading 4", 11, NAVY)):
        st = d.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), "Calibri")
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if fonts.get(qn(attr)) is not None:
                del fonts.attrib[qn(attr)]

    new_numbering = _docx_restart_numbering(d)
    shift = 1 if any(b["type"] == "title" for b in blocks) else 0  # "## Section" is top level under a "# Title"
    first_content = True
    for b in blocks:
        t = b["type"]
        if t == "title":
            d.add_paragraph(b["text"], style="Title")
        elif t == "heading":
            d.add_heading(b["text"], level=max(1, b["level"] - shift))
        elif t == "paragraph":
            _docx_runs(d.add_paragraph(), b["text"])
        elif t == "bullets":
            for text, level in b["items"]:
                p = d.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
                _docx_runs(p, text)
        elif t == "numbered":
            try:
                num_id = new_numbering()
            except Exception:
                num_id = None
            for text, level in b["items"]:
                p = d.add_paragraph(style="List Number")
                if num_id is not None:
                    num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
                    num_pr.get_or_add_ilvl().val = level
                    num_pr.get_or_add_numId().val = num_id
                _docx_runs(p, text)
        elif t == "table":
            rows = [b["header"]] + b["rows"]
            tb = d.add_table(rows=len(rows), cols=len(b["header"]))
            tb.style = "Table Grid"
            tb.autofit = True
            for ri, row in enumerate(rows):
                for ci, text in enumerate(row):
                    cell = tb.cell(ri, ci)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    if ri == 0:
                        _docx_runs(p, text, size=10.5, color="FFFFFF", force_bold=True)
                        tc_pr = cell._tc.get_or_add_tcPr()
                        shd = OxmlElement("w:shd")
                        shd.set(qn("w:val"), "clear")
                        shd.set(qn("w:color"), "auto")
                        shd.set(qn("w:fill"), NAVY)
                        tc_pr.append(shd)
                    else:
                        _docx_runs(p, text, size=10.5)
            hdr = tb.rows[0]._tr.get_or_add_trPr()
            flag = OxmlElement("w:tblHeader")
            flag.set(qn("w:val"), "true")
            hdr.append(flag)
            d.add_paragraph().paragraph_format.space_after = Pt(4)
        elif t == "pagebreak":
            if not first_content:
                d.add_page_break()
        first_content = False
    cp = d.core_properties
    cp.title = plain(title or document_title(blocks, "Document"))
    cp.author = AUTHOR
    d.save(path)
    return None


# ---------------------------------------------------------------- PowerPoint
def _blocks_to_slides(blocks):
    """Blocks -> slide dicts: {"kind": title|content|table, "title", "subtitle", "bullets": [(text, level)], "table"}."""
    slides, cur = [], None

    def new(kind, title):
        s = {"kind": kind, "title": title, "subtitle": "", "bullets": [], "table": None}
        slides.append(s)
        return s

    for b in blocks:
        t = b["type"]
        if t == "title":
            cur = new("title", b["text"])
        elif t == "heading":
            if b["level"] <= 2 or cur is None:
                cur = new("content", b["text"])
            else:
                cur["bullets"].append(("**%s**" % b["text"], 0, False))
        elif t == "paragraph":
            if cur is None:
                cur = new("content", "Overview")
            if cur["kind"] == "title" and not cur["subtitle"] and not cur["bullets"]:
                cur["subtitle"] = b["text"]
            else:
                if cur["kind"] == "title":
                    cur = new("content", cur["title"])
                cur["bullets"].append((b["text"], 0, False))
        elif t in ("bullets", "numbered"):
            if cur is None or cur["kind"] == "title":
                cur = new("content", cur["title"] if cur else "Overview")
            for text, level in b["items"]:
                cur["bullets"].append((text, level, t == "numbered"))
        elif t == "table":
            if cur is None or cur["kind"] == "title" or cur["table"] or len(cur["bullets"]) > 3:
                cur = new("table", cur["title"] if cur else "Table")
            cur["table"] = b
            cur["kind"] = "table"
    out = []
    for s in slides:  # split overflowing slides
        if s["table"]:
            rows = s["table"]["rows"]
            per = 7 if not s["bullets"] else 5
            chunks = [rows[i:i + per] for i in range(0, len(rows), per)] or [[]]
            for k, ch in enumerate(chunks):
                c = dict(s)
                c["title"] = s["title"] + (" (cont.)" if k else "")
                c["table"] = {"type": "table", "header": s["table"]["header"], "rows": ch}
                c["bullets"] = s["bullets"] if k == 0 else []
                out.append(c)
            continue
        weight = lambda item: 1 + len(plain(item[0])) // 75
        chunks, cur_chunk, w = [], [], 0
        for item in s["bullets"]:
            if cur_chunk and w + weight(item) > 9:
                chunks.append(cur_chunk)
                cur_chunk, w = [], 0
            cur_chunk.append(item)
            w += weight(item)
        chunks.append(cur_chunk)
        for k, ch in enumerate(chunks):
            c = dict(s)
            c["title"] = s["title"] + (" (cont.)" if k else "")
            c["bullets"] = ch
            out.append(c)
    return out


def create_pptx(path, blocks, title="", allow_formulas=False, warnings=None, slides_requested=None):
    warnings = warnings if warnings is not None else []
    import pptx
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt

    slides = _blocks_to_slides(blocks)
    if not slides:
        raise DocError("The presentation has no slides.")
    if len(slides) > MAX_SLIDES:
        warnings.append("The presentation was cut to %d slides." % MAX_SLIDES)
        slides = slides[:MAX_SLIDES]
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    rgb = RGBColor.from_string

    def add_runs(par, text, size, color, bold=False):
        for t, b, i in parse_inline(text):
            r = par.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = True if (b or bold) else None
            r.font.italic = True if i else None
            r.font.color.rgb = rgb(color)

    def place(shape, l, t, w, h):
        shape.left, shape.top, shape.width, shape.height = Inches(l), Inches(t), Inches(w), Inches(h)

    def drop(shape):
        shape._element.getparent().remove(shape._element)

    for s in slides:
        if s["kind"] == "title":
            sl = prs.slides.add_slide(prs.slide_layouts[0])
            sl.background.fill.solid()
            sl.background.fill.fore_color.rgb = rgb(NAVY)
            ttl = sl.shapes.title
            place(ttl, 0.9, 2.2, 11.5, 1.8)
            ttl.text_frame.word_wrap = True
            ttl.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
            par = ttl.text_frame.paragraphs[0]
            par.alignment = 1
            add_runs(par, plain(s["title"]), 48 if len(s["title"]) <= 36 else 38, "FFFFFF", bold=True)
            sub = sl.placeholders[1]
            if s["subtitle"]:
                place(sub, 0.9, 4.25, 11.5, 1.3)
                sub.text_frame.word_wrap = True
                p2 = sub.text_frame.paragraphs[0]
                p2.alignment = 1
                add_runs(p2, plain(s["subtitle"]), 26, "C9D8EE")
            else:
                drop(sub)
            bar = sl.shapes.add_shape(1, Inches(0.9), Inches(4.1), Inches(1.6), Inches(0.07))
            bar.fill.solid()
            bar.fill.fore_color.rgb = rgb("5BA3E0")
            bar.line.fill.background()
            continue
        sl = prs.slides.add_slide(prs.slide_layouts[1 if s["kind"] == "content" else 5])
        ttl = sl.shapes.title
        place(ttl, 0.7, 0.4, 11.9, 1.0)
        ttl.text_frame.word_wrap = True
        ttl.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        tp = ttl.text_frame.paragraphs[0]
        tp.alignment = 1
        n = len(s["title"])
        add_runs(tp, plain(s["title"]), 38 if n <= 40 else 30 if n <= 65 else 24, NAVY, bold=True)
        bar = sl.shapes.add_shape(1, Inches(0.7), Inches(1.45), Inches(1.2), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(ACCENT)
        bar.line.fill.background()
        top = 1.8
        if s["kind"] == "content":
            body = sl.placeholders[1]
            place(body, 0.7, top, 11.9, 5.0)
            tf = body.text_frame
            tf.word_wrap = True
            weight = sum(1 + len(plain(item[0])) // 75 for item in s["bullets"])
            size = 30 if weight <= 5 else 26 if weight <= 7 else 22
            first = True
            for text, level, numbered in s["bullets"]:
                par = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                par.level = min(level, 2)
                par.space_after = Pt(10)
                if numbered:
                    ppr = par._p.get_or_add_pPr()
                    ppr.set("marL", str(Inches(0.5 + 0.4 * min(level, 2))))
                    ppr.set("indent", str(-Inches(0.45)))
                    auto = ppr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"})
                    ppr.append(auto)
                add_runs(par, text, size - 3 * min(level, 2), "222222")
            if first:
                drop(body)
        else:
            if s["bullets"]:
                box = sl.shapes.add_textbox(Inches(0.7), Inches(top), Inches(11.9), Inches(1.2))
                box.text_frame.word_wrap = True
                for k, (text, level, _num) in enumerate(s["bullets"]):
                    par = box.text_frame.paragraphs[0] if k == 0 else box.text_frame.add_paragraph()
                    add_runs(par, text, 20, "222222")
                top += 0.4 + 0.4 * len(s["bullets"])
            tbl = s["table"]
            rows = [tbl["header"]] + tbl["rows"]
            ncol = len(tbl["header"])
            height = min(5.2, 0.5 * len(rows))
            shape = sl.shapes.add_table(len(rows), ncol, Inches(0.7), Inches(top), Inches(11.9), Inches(height))
            grid = shape.table
            lens = [min(max([len(plain(r[c])) for r in rows] + [3]), 40) for c in range(ncol)]
            total = float(sum(max(l, 6) for l in lens))
            for c in range(ncol):
                grid.columns[c].width = Emu(int(Inches(11.9) * max(lens[c], 6) / total))
            for ri, row in enumerate(rows):
                for ci, text in enumerate(row):
                    cell = grid.cell(ri, ci)
                    cell.text_frame.word_wrap = True
                    par = cell.text_frame.paragraphs[0]
                    if ri == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb(NAVY)
                        add_runs(par, text, 18, "FFFFFF", bold=True)
                    else:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb(LIGHT if ri % 2 == 0 else "FFFFFF")
                        add_runs(par, text, 17, "222222")
    cp = prs.core_properties
    cp.title = plain(title or document_title(blocks, "Presentation"))
    cp.author = AUTHOR
    prs.save(path)
    if slides_requested and slides_requested != len(slides):
        warnings.append("You asked for %d slides but the content made %d." % (slides_requested, len(slides)))
    return {"slides": len(slides)}


# ---------------------------------------------------------------- Excel
_NUM_RE = re.compile(r"^-?(?:0|[1-9]\d{0,14})(?:\.\d{1,10})?$")
_NUM_COMMA_RE = re.compile(r"^-?[1-9]\d{0,2}(?:,\d{3})+(?:\.\d+)?$")
_PCT_RE = re.compile(r"^-?\d{1,6}(?:\.\d{1,4})?%$")
_REF_RE = re.compile(r"^\$?[A-Za-z]{1,3}\$?\d{1,7}$")
_FORMULA_CHARS_RE = re.compile(r'^[A-Za-z0-9_$.,:+\-*/^()<>=&%"\s]+$')
SAFE_FUNCS = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "ROUND", "ABS", "IF", "PRODUCT", "MEDIAN"}


def safe_formula(text):
    """True only for simple, same-sheet arithmetic formulas (no links, no external data, no unknown functions)."""
    if not text.startswith("=") or len(text) > 300:
        return False
    body = re.sub(r'"[^"\n]*"', '""', text)
    if not _FORMULA_CHARS_RE.match(body) or body.count("(") != body.count(")"):
        return False
    for m in re.finditer(r"[A-Za-z_][A-Za-z0-9_.$]*", body):
        word, end = m.group(0), m.end()
        if body[end:end + 1] == "(":
            if word.upper() not in SAFE_FUNCS:
                return False
        elif not (_REF_RE.match(word) or word.upper() in ("TRUE", "FALSE")):
            return False
    return True


def _cell_value(text, allow_formulas, excel_row, warnings):
    """(value, number_format, force_text). Numbers become numbers; '=...' never becomes an accidental formula."""
    s = text.strip()
    if s.startswith("="):
        f = s.replace("{row}", str(excel_row))
        if allow_formulas and safe_formula(f):
            return f, None, False
        warnings.append("'%s' was kept as text (formulas are only added when you ask for them, and only simple ones)."
                        % s[:40])
        return s, None, True
    if _NUM_RE.match(s):
        return (float(s) if "." in s else int(s)), None, False
    if _NUM_COMMA_RE.match(s):
        v = s.replace(",", "")
        return (float(v) if "." in v else int(v)), "#,##0.00" if "." in v else "#,##0", False
    if _PCT_RE.match(s):
        dec = len(s.rstrip("%").partition(".")[2])
        return float(s.rstrip("%")) / 100.0, "0%" if not dec else "0." + "0" * dec + "%", False
    return plain(text) if s else None, None, False


def _sheet_name(raw, used):
    name = re.sub(r"[\[\]:*?/\\]", " ", plain(raw or "")).strip().strip("'")
    name = re.sub(r"\s+", " ", name)[:31].strip().strip("'") or "Sheet"
    base, k = name, 2
    while name.lower() in used:
        suffix = " %d" % k
        name = base[:31 - len(suffix)] + suffix
        k += 1
    used.add(name.lower())
    return name


def create_xlsx(path, blocks, title="", allow_formulas=False, warnings=None):
    warnings = warnings if warnings is not None else []
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sheets, heading, ignored = [], "", 0
    for b in blocks:
        if b["type"] in ("title", "heading"):
            heading = b["text"]
        elif b["type"] == "table":
            sheets.append((heading, b))
            heading = ""
        elif b["type"] in ("paragraph", "bullets", "numbered"):
            ignored += 1
    if not sheets:
        raise DocError("A spreadsheet needs a table (rows and columns), but the content has none.")
    if ignored:
        warnings.append("Text outside the tables was not added to the workbook.")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used = set()
    thin = Side(style="thin", color="B8C4D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
    for heading, tb in sheets[:200]:
        ws = wb.create_sheet(_sheet_name(heading or ("Sheet%d" % (len(used) + 1)), used))
        header = [plain(h) for h in tb["header"]]
        ws.append(header)
        widths = [len(h) for h in header]
        for ri, row in enumerate(tb["rows"], start=2):
            cells = []
            for ci, text in enumerate(row):
                v, fmt, force = _cell_value(text, allow_formulas, ri, warnings)
                cells.append((v, fmt, force))
                widths[ci] = max(widths[ci], 10 if isinstance(v, str) and v.startswith("=") and not force else len(plain(text)))
            ws.append([c[0] for c in cells])
            for ci, (v, fmt, force) in enumerate(cells, start=1):
                cell = ws.cell(row=ri, column=ci)
                if force:
                    cell.data_type = "s"  # stored as text: Excel will never evaluate it
                if fmt:
                    cell.number_format = fmt
        for ci in range(1, len(header) + 1):
            c = ws.cell(row=1, column=ci)
            c.font, c.fill, c.border = head_font, head_fill, border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(ci)].width = max(10, min(60, widths[ci - 1] * 1.15 + 3))
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(header)):
            for c in row:
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 50)
        ws.row_dimensions[1].height = 22
        ws.freeze_panes = "A2"
        if tb["rows"]:
            ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(header)), ws.max_row)
    if len(sheets) > 200:
        warnings.append("Only the first 200 tables were saved.")
    wb.properties.title = plain(title or document_title(blocks, "Workbook"))
    wb.properties.creator = AUTHOR
    wb.save(path)
    return {"sheets": list(wb.sheetnames)}


_BUILDERS.update({"pdf": create_pdf, "docx": create_docx, "pptx": create_pptx, "xlsx": create_xlsx})
